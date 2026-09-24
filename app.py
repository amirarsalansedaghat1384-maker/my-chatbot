"""
چت‌بات فارسی با Flask و OpenRouter
قابلیت‌ها: چند مکالمه، ذخیره دائمی، جدا بودن کاربران، آپلود عکس،
محدودیت پیام روزانه، پاسخ استریم (تایپ‌شونده)، جستجو در تاریخچه
"""

from flask import Flask, render_template, request, jsonify, g, Response
import requests
import os
import json
import uuid
import time
import datetime
import io
import re
from pypdf import PdfReader
from docx import Document
import openpyxl

app = Flask(__name__)

API_KEY = os.environ.get("OPENROUTER_API_KEY")

API_URL = "https://openrouter.ai/api/v1/chat/completions"
MODEL = "openrouter/free"

SYSTEM_PROMPT = (
    "تو یک دستیار هوش مصنوعی فارسی‌زبان هستی. همیشه فقط و فقط به زبان فارسی روان و "
    "طبیعی پاسخ بده و هیچ کلمه یا جمله‌ای از زبان‌های دیگر (انگلیسی، اسپانیایی، رومانیایی و غیره) "
    "قاطی جواب‌هات نکن، مگر اینکه کاربر مستقیم از تو بخواد یا اسم خاص/فنی باشه که معادل فارسی نداره. "
    "برای فرمت‌بندی جواب‌هات از Markdown استاندارد استفاده کن: **متن پررنگ** برای تاکید، "
    "لیست‌ها با - یا شماره، و تیترها در صورت نیاز. جواب‌هات رو تمیز، خوانا و مرتب بنویس."
)

HISTORY_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "chat_history.json")

COOKIE_NAME = "user_id"
COOKIE_MAX_AGE = 60 * 60 * 24 * 365 * 2  # دو سال

MAX_DAILY_MESSAGES = 30  # محدودیت پیام رایگان روزانه برای هر کاربر
MAX_DOC_CHARS = 12000  # حداکثر تعداد کاراکتری که از یه فایل استخراج می‌کنیم
MAX_UPLOAD_SIZE = 8 * 1024 * 1024  # حداکثر حجم فایل آپلودی: ۸ مگابایت

MEMORY_TRIGGERS = ["یادت باشه", "به خاطر بسپار", "فراموش نکن", "یادت بمونه", "حفظ کن که"]
MAX_MEMORY_NOTES = 30  # حداکثر تعداد یادداشت حافظه برای هر کاربر


def load_data():
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if "users" in data:
                    return data
        except (json.JSONDecodeError, IOError):
            pass
    return {"users": {}}


def save_data(data):
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


data_store = load_data()


def get_user_entry():
    user_id = g.user_id
    if user_id not in data_store["users"]:
        data_store["users"][user_id] = {"chats": {}}
    return data_store["users"][user_id]


def get_user_chats():
    return get_user_entry()["chats"]


def get_user_memory():
    entry = get_user_entry()
    if "memory" not in entry:
        entry["memory"] = []
    return entry["memory"]


def check_memory_trigger(text):
    """اگه پیام کاربر شامل یکی از عبارات یادآوری بود، اون رو به‌عنوان یادداشت حافظه برمی‌گردونه"""
    for trigger in MEMORY_TRIGGERS:
        if trigger in text:
            return text.strip()
    return None


def add_memory_note(note):
    notes = get_user_memory()
    notes.append({"text": note, "created": time.time()})
    if len(notes) > MAX_MEMORY_NOTES:
        del notes[0]
    save_data(data_store)


def build_memory_context():
    notes = get_user_memory()
    if not notes:
        return ""
    lines = "\n".join(f"- {n['text']}" for n in notes)
    return (
        "\n\nچیزهایی که کاربر قبلاً ازت خواسته به خاطر بسپاری (در مکالمات قبلی):\n"
        f"{lines}\n"
        "این‌ها رو در نظر بگیر، ولی فقط وقتی مرتبطه ازشون استفاده کن."
    )


def extract_text_from_file(filename, file_bytes):
    """متن رو از فایل PDF/Word/Excel/متنی استخراج می‌کنه"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(file_bytes))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == "docx":
            doc = Document(io.BytesIO(file_bytes))
            text = "\n".join(p.text for p in doc.paragraphs)
        elif ext in ("xlsx", "xlsm"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"--- شیت: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    line = " | ".join(str(cell) for cell in row if cell is not None)
                    if line.strip():
                        parts.append(line)
            text = "\n".join(parts)
        elif ext in ("txt", "csv", "md"):
            text = file_bytes.decode("utf-8", errors="ignore")
        elif ext == "doc":
            return None, "فایل‌های Word با فرمت قدیمی (.doc) پشتیبانی نمی‌شن - لطفاً به .docx تبدیلش کن."
        elif ext == "xls":
            return None, "فایل‌های Excel با فرمت قدیمی (.xls) پشتیبانی نمی‌شن - لطفاً به .xlsx تبدیلش کن."
        else:
            return None, f"فرمت .{ext} پشتیبانی نمی‌شه."
    except Exception as e:
        return None, f"خطا در خوندن فایل: {str(e)}"

    text = text.strip()
    if not text:
        return None, "متنی از این فایل استخراج نشد (ممکنه اسکن‌شده یا خالی باشه)."

    truncated = len(text) > MAX_DOC_CHARS
    if truncated:
        text = text[:MAX_DOC_CHARS]

    return text, ("truncated" if truncated else None)


def make_title(first_message):
    title = first_message.strip()
    if len(title) > 30:
        title = title[:30] + "..."
    return title or "مکالمه جدید"


def check_and_increment_usage():
    """برمی‌گردونه: (مجاز است؟, پیام خطا در صورت رد شدن)"""
    user_entry = get_user_entry()
    today = datetime.date.today().isoformat()
    usage = user_entry.setdefault("usage", {"date": today, "count": 0})
    if usage["date"] != today:
        usage["date"] = today
        usage["count"] = 0
    if usage["count"] >= MAX_DAILY_MESSAGES:
        return False, f"محدودیت {MAX_DAILY_MESSAGES} پیام رایگان امروزت تموم شد. فردا دوباره امتحان کن."
    usage["count"] += 1
    return True, None


def build_content(user_message, image_data):
    if image_data:
        text_part = user_message or "این تصویر رو توضیح بده"
        return [
            {"type": "text", "text": text_part},
            {"type": "image_url", "image_url": {"url": image_data}},
        ]
    return user_message


def prepare_chat(chat_id, user_chats, user_message):
    if not chat_id or chat_id not in user_chats:
        chat_id = str(uuid.uuid4())
        user_chats[chat_id] = {
            "title": "مکالمه جدید",
            "messages": [],
            "created": time.time(),
        }
    chat_obj = user_chats[chat_id]
    if len(chat_obj["messages"]) == 0:
        chat_obj["title"] = make_title(user_message or "تصویر")
    return chat_id, chat_obj


@app.before_request
def ensure_user_id():
    user_id = request.cookies.get(COOKIE_NAME)
    if not user_id:
        user_id = str(uuid.uuid4())
        g.new_user_id = user_id
    g.user_id = user_id


@app.after_request
def set_user_cookie(response):
    if hasattr(g, "new_user_id"):
        response.set_cookie(
            COOKIE_NAME,
            g.new_user_id,
            max_age=COOKIE_MAX_AGE,
            httponly=True,
            samesite="Lax",
        )
    return response


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/usage", methods=["GET"])
def usage_status():
    user_entry = get_user_entry()
    today = datetime.date.today().isoformat()
    usage = user_entry.get("usage", {"date": today, "count": 0})
    if usage["date"] != today:
        usage = {"date": today, "count": 0}
    return jsonify({"used": usage["count"], "limit": MAX_DAILY_MESSAGES})


@app.route("/upload-document", methods=["POST"])
def upload_document():
    if "file" not in request.files:
        return jsonify({"error": "فایلی ارسال نشد"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "فایلی انتخاب نشد"}), 400

    file_bytes = file.read()
    if len(file_bytes) > MAX_UPLOAD_SIZE:
        return jsonify({"error": "حجم فایل بیشتر از ۸ مگابایته"}), 400

    text, note = extract_text_from_file(file.filename, file_bytes)
    if text is None:
        return jsonify({"error": note}), 400

    return jsonify({
        "filename": file.filename,
        "text": text,
        "truncated": note == "truncated",
    })


@app.route("/memory", methods=["GET"])
def get_memory():
    return jsonify({"notes": get_user_memory()})


@app.route("/memory", methods=["POST"])
def add_memory_manual():
    note = (request.json or {}).get("text", "").strip()
    if not note:
        return jsonify({"error": "متن نمی‌تونه خالی باشه"}), 400
    add_memory_note(note)
    return jsonify({"status": "ok", "notes": get_user_memory()})


@app.route("/memory/<int:index>", methods=["DELETE"])
def delete_memory(index):
    notes = get_user_memory()
    if 0 <= index < len(notes):
        del notes[index]
        save_data(data_store)
    return jsonify({"notes": get_user_memory()})


@app.route("/chats", methods=["GET"])
def list_chats():
    user_chats = get_user_chats()
    chats = [
        {"id": cid, "title": c["title"], "created": c["created"]}
        for cid, c in user_chats.items()
    ]
    chats.sort(key=lambda c: c["created"], reverse=True)
    return jsonify({"chats": chats})


@app.route("/chats/search", methods=["GET"])
def search_chats():
    query = request.args.get("q", "").strip()
    user_chats = get_user_chats()
    if not query:
        return jsonify({"chats": []})

    results = []
    for cid, c in user_chats.items():
        found = query in c["title"]
        if not found:
            for m in c["messages"]:
                content = m["content"]
                if isinstance(content, str):
                    text = content
                else:
                    text = " ".join(p.get("text", "") for p in content if p.get("type") == "text")
                if query in text:
                    found = True
                    break
        if found:
            results.append({"id": cid, "title": c["title"], "created": c["created"]})

    results.sort(key=lambda c: c["created"], reverse=True)
    return jsonify({"chats": results})


@app.route("/chats/new", methods=["POST"])
def new_chat():
    user_chats = get_user_chats()
    chat_id = str(uuid.uuid4())
    user_chats[chat_id] = {
        "title": "مکالمه جدید",
        "messages": [],
        "created": time.time(),
    }
    save_data(data_store)
    return jsonify({"chat_id": chat_id})


@app.route("/chats/<chat_id>", methods=["GET"])
def get_chat(chat_id):
    user_chats = get_user_chats()
    chat = user_chats.get(chat_id)
    if not chat:
        return jsonify({"error": "مکالمه پیدا نشد"}), 404
    return jsonify({"messages": chat["messages"]})


@app.route("/chats/<chat_id>", methods=["DELETE"])
def delete_chat(chat_id):
    user_chats = get_user_chats()
    if chat_id in user_chats:
        del user_chats[chat_id]
        save_data(data_store)
    return jsonify({"status": "حذف شد"})


@app.route("/chats/<chat_id>/rename", methods=["POST"])
def rename_chat(chat_id):
    user_chats = get_user_chats()
    new_title = (request.json or {}).get("title", "").strip()
    if not new_title:
        return jsonify({"error": "اسم نمی‌تونه خالی باشه"}), 400
    if chat_id not in user_chats:
        return jsonify({"error": "مکالمه پیدا نشد"}), 404
    user_chats[chat_id]["title"] = new_title[:50]
    save_data(data_store)
    return jsonify({"status": "ok", "title": user_chats[chat_id]["title"]})


@app.route("/chat", methods=["POST"])
def chat():
    """نسخه غیر-استریم (پشتیبان/جایگزین)"""
    user_chats = get_user_chats()
    chat_id = request.json.get("chat_id")
    user_message = request.json.get("message", "").strip()
    image_data = request.json.get("image")

    if not user_message and not image_data:
        return jsonify({"error": "پیام خالی است"}), 400

    allowed, err_msg = check_and_increment_usage()
    if not allowed:
        return jsonify({"error": err_msg}), 429

    memory_note = check_memory_trigger(user_message)
    if memory_note:
        add_memory_note(memory_note)

    chat_id, chat_obj = prepare_chat(chat_id, user_chats, user_message)
    content = build_content(user_message, image_data)
    chat_obj["messages"].append({"role": "user", "content": content})
    save_data(data_store)

    try:
        headers = {"Authorization": f"Bearer {API_KEY}", "content-type": "application/json"}
        payload = {
            "model": MODEL,
            "messages": [{"role": "system", "content": SYSTEM_PROMPT + build_memory_context()}] + chat_obj["messages"],
        }
        res = requests.post(API_URL, headers=headers, json=payload, timeout=60)
        response_data = res.json()

        if res.status_code != 200:
            error_msg = response_data.get("error", {}).get("message", str(response_data))
            return jsonify({"error": f"خطا در ارتباط با API: {error_msg}"}), 500

        assistant_reply = response_data["choices"][0]["message"]["content"]
        chat_obj["messages"].append({"role": "assistant", "content": assistant_reply})
        save_data(data_store)

        return jsonify({"reply": assistant_reply, "chat_id": chat_id})

    except Exception as e:
        return jsonify({"error": f"خطا در ارتباط با API: {str(e)}"}), 500


@app.route("/chat/stream", methods=["POST"])
def chat_stream():
    """نسخه استریم - جواب رو به‌صورت تکه‌تکه (تایپ‌شونده) برمی‌گردونه"""
    user_chats = get_user_chats()
    chat_id = request.json.get("chat_id")
    user_message = request.json.get("message", "").strip()
    image_data = request.json.get("image")

    if not user_message and not image_data:
        return jsonify({"error": "پیام خالی است"}), 400

    allowed, err_msg = check_and_increment_usage()
    if not allowed:
        return jsonify({"error": err_msg}), 429

    memory_note = check_memory_trigger(user_message)
    if memory_note:
        add_memory_note(memory_note)

    chat_id, chat_obj = prepare_chat(chat_id, user_chats, user_message)
    content = build_content(user_message, image_data)
    chat_obj["messages"].append({"role": "user", "content": content})
    save_data(data_store)

    system_content = SYSTEM_PROMPT + build_memory_context()

    def generate():
        full_reply = ""
        try:
            headers = {"Authorization": f"Bearer {API_KEY}", "content-type": "application/json"}
            payload = {
                "model": MODEL,
                "messages": [{"role": "system", "content": system_content}] + chat_obj["messages"],
                "stream": True,
            }
            with requests.post(API_URL, headers=headers, json=payload, timeout=120, stream=True) as res:
                if res.status_code != 200:
                    yield f"data: {json.dumps({'error': 'خطا در ارتباط با سرور هوش مصنوعی'})}\n\n"
                    return
                for line in res.iter_lines():
                    if not line:
                        continue
                    decoded = line.decode("utf-8", errors="ignore")
                    if not decoded.startswith("data: "):
                        continue
                    data_str = decoded[6:]
                    if data_str.strip() == "[DONE]":
                        break
                    try:
                        chunk = json.loads(data_str)
                        delta = chunk["choices"][0]["delta"].get("content", "")
                        if delta:
                            full_reply += delta
                            yield f"data: {json.dumps({'delta': delta})}\n\n"
                    except Exception:
                        continue
        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
        finally:
            if full_reply:
                chat_obj["messages"].append({"role": "assistant", "content": full_reply})
                save_data(data_store)
            yield f"data: {json.dumps({'done': True, 'chat_id': chat_id})}\n\n"

    return Response(generate(), mimetype="text/event-stream")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=True, host="0.0.0.0", port=port)
